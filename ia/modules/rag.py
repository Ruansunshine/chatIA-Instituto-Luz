"""RAG: ingestao (PDF/XLSX -> chunks -> ChromaDB) e busca de contexto pro
/responder. Isso tudo e' CODIGO deterministico (parsing de arquivo, busca
vetorial) - ver criterio.py. So o texto final devolvido ao usuario passa
pelo Ollama, com o contexto encontrado aqui injetado no prompt.
"""

import os
from pathlib import Path

import chromadb
from openpyxl import load_workbook
from pypdf import PdfReader

CHROMA_HOST = os.getenv("CHROMA_HOST", "chroma")
CHROMA_PORT = int(os.getenv("CHROMA_PORT", "8000"))
NOME_COLECAO = "instituto_luz"

PASTA_DADOS = Path(__file__).resolve().parent.parent / "dados"
CAMINHO_PDF = PASTA_DADOS / "guia_do_paciente.pdf"
CAMINHO_XLSX = PASTA_DADOS / "precos_profissionais.xlsx"

# Chunking generico por janela de caracteres - NAO depende de conhecer o
# conteudo do PDF de antemao (nada de titulo/secao hardcoded). O sistema tem
# que ler o documento de verdade, nao ter o conteudo dele "entregue" no
# codigo - isso teria que funcionar em qualquer PDF, nao so' neste.
TAMANHO_CHUNK = 800  # caracteres
SOBREPOSICAO_CHUNK = 150  # caracteres repetidos entre chunks vizinhos, pra
# nao cortar uma frase relevante bem na fronteira de dois chunks


# Client unico reaproveitado entre chamadas - antes cada chamada criava um
# chromadb.HttpClient() novo (com seu proprio pool de conexao HTTP) e nunca
# fechava o anterior. Achado na pratica via auditoria: 10 conexoes TCP
# abertas e ociosas do container Python pro Chroma, acumulando a cada
# ingestao/busca.
_CLIENTE = None


def _cliente_chroma() -> chromadb.HttpClient:
    global _CLIENTE
    if _CLIENTE is None:
        _CLIENTE = chromadb.HttpClient(host=CHROMA_HOST, port=CHROMA_PORT)
    return _CLIENTE


def gerar_chunks_pdf() -> list[dict]:
    """Le o PDF de verdade (pypdf.extract_text por pagina) e corta em janelas
    de tamanho fixo com sobreposicao - funciona pra qualquer PDF, sem saber
    nada do conteudo/estrutura dele antes de abrir o arquivo."""
    leitor = PdfReader(str(CAMINHO_PDF))
    texto_bruto = "\n".join(pagina.extract_text() or "" for pagina in leitor.pages)
    texto = " ".join(texto_bruto.split())  # normaliza espacos e quebras de linha

    chunks = []
    inicio = 0
    numero = 0
    passo = TAMANHO_CHUNK - SOBREPOSICAO_CHUNK
    while inicio < len(texto):
        fim = min(inicio + TAMANHO_CHUNK, len(texto))
        trecho = texto[inicio:fim].strip()
        if trecho:
            numero += 1
            chunks.append(
                {
                    "texto": trecho,
                    "fonte": "guia_do_paciente.pdf",
                    "rotulo": f"Guia do Paciente (trecho {numero})",
                    "tipo": "guia",
                }
            )
        if fim == len(texto):
            break
        inicio += passo
    return chunks


def gerar_chunks_planilha() -> list[dict]:
    planilha = load_workbook(str(CAMINHO_XLSX), data_only=True)
    chunks = []

    aba_precos = planilha["Precos_Pacotes"]
    por_especialidade: dict[str, list[str]] = {}
    for linha in aba_precos.iter_rows(min_row=2, values_only=True):
        if not linha or len(linha) < 6:
            continue
        especialidade, tipo, preco_unit, preco_total, validade, obs = linha[:6]
        # pula linha incompleta (ex: nota de rodape na planilha, que so tem
        # texto na 1a coluna) - so processa linha com preco de verdade.
        if None in (especialidade, tipo, preco_unit, preco_total):
            continue
        por_especialidade.setdefault(especialidade, []).append(
            f"{tipo}: R$ {preco_unit:.2f} (total R$ {preco_total:.2f}, "
            f"validade {validade} dias) - {obs}"
        )
    for especialidade, itens in por_especialidade.items():
        chunks.append(
            {
                "texto": f"Preços da especialidade {especialidade}:\n" + "\n".join(itens),
                "fonte": "precos_profissionais.xlsx",
                "rotulo": f"Planilha de preços, {especialidade}",
                "tipo": "preco",
            }
        )

    aba_profissionais = planilha["Profissionais_Horarios"]
    for linha in aba_profissionais.iter_rows(min_row=2, values_only=True):
        if not linha or len(linha) < 6:
            continue
        nome, especialidade, dias, inicio, fim, anos = linha[:6]
        if None in (nome, especialidade, dias, inicio, fim):
            continue
        chunks.append(
            {
                "texto": (
                    f"{nome} atende a especialidade {especialidade}, nos dias "
                    f"{dias}, das {inicio} às {fim}. {anos} anos de experiência."
                ),
                "fonte": "precos_profissionais.xlsx",
                "rotulo": f"Planilha de horários, {nome}",
                "tipo": "profissional",
            }
        )

    return chunks


def ingerir() -> int:
    """Idempotente: apaga a colecao e recria do zero - roda no startup do
    container (ver main.py), nao pode duplicar chunk a cada restart."""
    cliente = _cliente_chroma()
    try:
        cliente.delete_collection(NOME_COLECAO)
    except Exception:
        pass
    colecao = cliente.create_collection(NOME_COLECAO)

    chunks = gerar_chunks_pdf() + gerar_chunks_planilha()
    if not chunks:
        return 0

    colecao.add(
        ids=[f"chunk-{i}" for i in range(len(chunks))],
        documents=[c["texto"] for c in chunks],
        metadatas=[
            {"fonte": c["fonte"], "rotulo": c["rotulo"], "tipo": c["tipo"]} for c in chunks
        ],
    )
    return len(chunks)


# top_k=4: testado na pratica e reduzido pra 2 antes pra ganhar velocidade -
# causou alucinacao de preco (o chunk certo de "Planilha de precos, Pilates"
# ficava em 3o lugar no ranking de similaridade e ficava de fora). Correcao
# de valor pro paciente pesa mais que alguns segundos a mais de resposta.
def buscar_contexto(pergunta: str, top_k: int = 4) -> tuple[list[str], list[dict]]:
    cliente = _cliente_chroma()
    colecao = cliente.get_collection(NOME_COLECAO)
    resultado = colecao.query(query_texts=[pergunta], n_results=top_k)
    documentos = resultado["documents"][0] if resultado["documents"] else []
    metadados = resultado["metadatas"][0] if resultado["metadatas"] else []
    return documentos, metadados


# Pra pergunta agregada ("quantos profissionais", "quais especialidades")
# a busca por similaridade (top_k) as vezes deixa de fora chunk relevante -
# o modelo so ve o que veio, nao "sabe" que faltou o resto (ja vimos isso
# alucinar contagem errada). Como a base e' pequena (poucos chunks por
# tipo), buscar TODOS de uma categoria e' barato e garante que o modelo
# recebe o conjunto completo pra contar/listar direito.
def buscar_todos_por_tipo(tipo: str) -> tuple[list[str], list[dict]]:
    cliente = _cliente_chroma()
    colecao = cliente.get_collection(NOME_COLECAO)
    resultado = colecao.get(where={"tipo": tipo})
    documentos = resultado["documents"] or []
    metadados = resultado["metadatas"] or []
    return documentos, metadados


def montar_texto_contexto(documentos: list[str]) -> str:
    return "\n\n---\n\n".join(documentos)


def formatar_fontes(metadados: list[dict], limite: int = 2) -> str:
    # metadados ja vem ordenado por relevancia (chunk mais proximo primeiro).
    # So cita os mais relevantes - contexto usa mais chunks (top_k=4) pra
    # nao alucinar dado, mas exibir TODOS como "fonte" numa resposta de uma
    # frase e' ruido (as vezes um chunk pouco relevante entra no contexto so
    # de reforco e nao tem nada a ver com o que a resposta realmente usou).
    rotulos_unicos: list[str] = []
    for metadado in metadados:
        rotulo = metadado.get("rotulo")
        if rotulo and rotulo not in rotulos_unicos:
            rotulos_unicos.append(rotulo)
        if len(rotulos_unicos) >= limite:
            break
    return ", ".join(rotulos_unicos)
